import csv
import io
import re
from typing import Any

import httpx
from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError
from sqlalchemy.orm import Session

from app.api.dependencies import get_current_user, require_permission
from app.core.config import settings
from app.db.session import get_db
from app.models.category import Category
from app.models.customer import Customer
from app.models.manufacturer import Manufacturer
from app.models.product import Product
from app.models.supplier import Supplier
from app.models.user import User
from app.schemas.master_data import CategoryCreate, CustomerCreate, ManufacturerCreate, ProductCreate
from app.schemas.purchase import SupplierCreate

router = APIRouter()
GSTIN_PATTERN = re.compile(r"^[0-9]{2}[A-Z]{5}[0-9]{4}[A-Z][1-9A-Z]Z[0-9A-Z]$")

TEMPLATES: dict[str, list[str]] = {
    "customers": [
        "customer_code", "customer_name", "business_name", "customer_type", "contact_person",
        "phone", "email", "billing_address", "shipping_address", "gstin", "drug_license_number",
        "drug_license_type", "state", "state_code", "pincode", "credit_limit", "credit_days",
        "opening_balance", "is_active",
    ],
    "suppliers": [
        "supplier_code", "supplier_name", "contact_person", "phone", "email", "address", "gstin",
        "credit_days", "credit_limit", "is_active",
    ],
    "categories": ["name", "description", "is_active"],
    "manufacturers": ["name", "address", "phone", "email", "gstin", "is_active"],
    "products": [
        "product_code", "product_name", "generic_name", "brand_name", "category_name",
        "manufacturer_name", "schedule_category", "dosage_form", "strength", "hsn_code", "barcode",
        "gst_rate", "unit", "pack_size", "default_mrp", "default_selling_price", "minimum_sale_rate",
        "reorder_level", "is_active",
    ],
}


def _clean(value: Any):
    if isinstance(value, str):
        value = value.strip()
        if not value:
            return None
        if value.lower() in {"true", "yes", "y"}:
            return True
        if value.lower() in {"false", "no", "n"}:
            return False
    return value


def _normalise_row(row: dict[str, Any]) -> dict[str, Any]:
    return {str(k).strip(): _clean(v) for k, v in row.items() if k is not None and str(k).strip()}


def _read_rows(upload: UploadFile, raw: bytes) -> list[dict[str, Any]]:
    name = (upload.filename or "").lower()
    if name.endswith(".xlsx"):
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet = workbook.active
        values = list(sheet.iter_rows(values_only=True))
        if not values:
            return []
        headers = [str(value).strip() if value is not None else "" for value in values[0]]
        return [_normalise_row(dict(zip(headers, row))) for row in values[1:] if any(value is not None and str(value).strip() for value in row)]
    if name.endswith(".csv"):
        text = raw.decode("utf-8-sig")
        return [_normalise_row(row) for row in csv.DictReader(io.StringIO(text)) if any(_clean(v) is not None for v in row.values())]
    raise HTTPException(status_code=400, detail="Upload an .xlsx or .csv file")


def _duplicate(entity: str, values: dict[str, Any], db: Session, company_id: int):
    if entity == "customers":
        return db.query(Customer).filter(Customer.company_id == company_id, Customer.customer_code == values["customer_code"]).first()
    if entity == "suppliers":
        return db.query(Supplier).filter(Supplier.company_id == company_id, Supplier.supplier_code == values["supplier_code"]).first()
    if entity == "categories":
        return db.query(Category).filter(Category.company_id == company_id, Category.name.ilike(values["name"])).first()
    if entity == "manufacturers":
        return db.query(Manufacturer).filter(Manufacturer.company_id == company_id, Manufacturer.name.ilike(values["name"])).first()
    if entity == "products":
        return db.query(Product).filter(Product.company_id == company_id, Product.product_code == values["product_code"]).first()
    return None


def _validate_row(entity: str, row: dict[str, Any], db: Session, company_id: int) -> tuple[dict[str, Any] | None, str | None]:
    try:
        if entity == "customers":
            values = CustomerCreate.model_validate(row).model_dump()
        elif entity == "suppliers":
            values = SupplierCreate.model_validate(row).model_dump()
        elif entity == "categories":
            values = CategoryCreate.model_validate(row).model_dump()
        elif entity == "manufacturers":
            values = ManufacturerCreate.model_validate(row).model_dump()
        elif entity == "products":
            category_name = str(row.pop("category_name", "") or "").strip()
            manufacturer_name = str(row.pop("manufacturer_name", "") or "").strip()
            category = db.query(Category).filter(Category.company_id == company_id, Category.name.ilike(category_name)).first()
            if not category:
                return None, f"Unknown category '{category_name}'. Create/import categories first."
            row["category_id"] = category.id
            if manufacturer_name:
                manufacturer = db.query(Manufacturer).filter(Manufacturer.company_id == company_id, Manufacturer.name.ilike(manufacturer_name)).first()
                if not manufacturer:
                    return None, f"Unknown manufacturer '{manufacturer_name}'. Create/import manufacturers first."
                row["manufacturer_id"] = manufacturer.id
            values = ProductCreate.model_validate(row).model_dump()
        else:
            return None, "Unsupported import type"
    except ValidationError as exc:
        first = exc.errors()[0]
        location = ".".join(str(part) for part in first.get("loc", []))
        return None, f"{location}: {first.get('msg', 'Invalid value')}"

    if _duplicate(entity, values, db, company_id):
        key = values.get("customer_code") or values.get("supplier_code") or values.get("product_code") or values.get("name")
        return None, f"Duplicate record already exists: {key}"
    return values, None


def _model_for(entity: str):
    return {
        "customers": Customer,
        "suppliers": Supplier,
        "categories": Category,
        "manufacturers": Manufacturer,
        "products": Product,
    }[entity]


@router.get("/gstin/{gstin}")
def verify_gstin(
    gstin: str,
    current_user: User = Depends(get_current_user),
):
    del current_user
    value = gstin.strip().upper()
    if not GSTIN_PATTERN.fullmatch(value):
        raise HTTPException(status_code=400, detail="Enter a valid 15-character GSTIN")
    if settings.GST_PROVIDER.lower() != "gstinapi":
        raise HTTPException(status_code=503, detail="Configured GST provider is not supported")
    if not settings.GSTINAPI_KEY:
        raise HTTPException(status_code=503, detail="GST verification is not configured. Add GSTINAPI_KEY on the server.")
    try:
        response = httpx.get(
            f"{settings.GSTINAPI_BASE_URL.rstrip('/')}/v1/gstin/{value}",
            headers={"x-api-key": settings.GSTINAPI_KEY}, timeout=12.0,
        )
    except httpx.RequestError:
        raise HTTPException(status_code=502, detail="GST verification provider is currently unreachable")
    try:
        payload = response.json()
    except ValueError:
        raise HTTPException(status_code=502, detail="GST verification provider returned an invalid response")
    if response.status_code >= 400 or not payload.get("success", response.status_code == 200):
        raise HTTPException(status_code=response.status_code if 400 <= response.status_code < 500 else 502, detail=payload.get("error") or "GST verification failed")
    data = payload.get("data") or payload
    return {
        "gstin": data.get("gstin") or value,
        "legal_name": data.get("legal_name"),
        "trade_name": data.get("trade_name"),
        "status": data.get("status"),
        "taxpayer_type": data.get("taxpayer_type"),
        "business_constitution": data.get("business_constitution"),
        "registration_date": data.get("registration_date"),
        "cancellation_date": data.get("cancellation_date"),
        "state_code": data.get("state_code"),
        "state_jurisdiction": data.get("state_jurisdiction"),
        "address": data.get("address"),
        "credits_remaining": payload.get("credits_remaining"),
        "provider": "gstinapi",
    }


@router.get("/import-template/{entity}")
def download_import_template(entity: str, current_user: User = Depends(get_current_user)):
    del current_user
    if entity not in TEMPLATES:
        raise HTTPException(status_code=404, detail="Unsupported template type")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = entity.title()
    sheet.append(TEMPLATES[entity])
    if entity == "customers":
        sheet.append(["CUST001", "Example Pharmacy", "Example Pharmacy", "PHARMACY", "", "", "", "", "", "", "", "", "Telangana", "36", "", 50000, 30, 0, True])
    elif entity == "suppliers":
        sheet.append(["SUP001", "Example Distributor", "", "", "", "", "", 30, "", True])
    elif entity == "categories":
        sheet.append(["Tablets", "Example category", True])
    elif entity == "manufacturers":
        sheet.append(["Example Pharma Ltd", "", "", "", "", True])
    elif entity == "products":
        sheet.append(["PRD001", "Example Tablet", "", "", "Tablets", "Example Pharma Ltd", "", "Tablet", "500 mg", "3004", "", 12, "Strip", 10, 100, 80, 75, 50, True])
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    headers = {"Content-Disposition": f'attachment; filename="medibill_{entity}_template.xlsx"'}
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=headers)


@router.post("/import/{entity}")
async def import_master_data(
    entity: str,
    preview: bool = Query(default=True),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("manage_inventory")),
):
    if entity not in TEMPLATES:
        raise HTTPException(status_code=404, detail="Unsupported import type")
    raw = await file.read()
    if len(raw) > settings.MAX_UPLOAD_SIZE_MB * 1024 * 1024:
        raise HTTPException(status_code=413, detail=f"File exceeds {settings.MAX_UPLOAD_SIZE_MB} MB limit")
    rows = _read_rows(file, raw)
    if not rows:
        raise HTTPException(status_code=400, detail="The uploaded file has no data rows")
    validated: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    seen: set[str] = set()
    for index, row in enumerate(rows, start=2):
        values, error = _validate_row(entity, dict(row), db, current_user.company_id)
        if values:
            identity = str(values.get("customer_code") or values.get("supplier_code") or values.get("product_code") or values.get("name") or "").lower()
            if identity in seen:
                error = f"Duplicate value appears more than once in file: {identity}"
                values = None
            else:
                seen.add(identity)
        if error:
            errors.append({"row": index, "error": error})
        else:
            validated.append(values or {})
    result = {"entity": entity, "total_rows": len(rows), "valid_rows": len(validated), "error_rows": len(errors), "errors": errors[:100], "preview": preview, "imported": 0}
    if preview or errors:
        return result
    model = _model_for(entity)
    try:
        for values in validated:
            db.add(model(company_id=current_user.company_id, **values))
        db.commit()
    except Exception:
        db.rollback()
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Import could not be committed. Check for duplicates or invalid references.")
    result["imported"] = len(validated)
    return result
